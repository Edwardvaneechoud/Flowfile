from pathlib import Path

from openpyxl import Workbook

from flowfile_core.flowfile.flow_data_engine.flow_data_engine import FlowDataEngine
from flowfile_core.flowfile.flow_graph import FlowGraph
from flowfile_core.flowfile.handler import FlowfileHandler
from flowfile_core.schemas import schemas
from flowfile_core.schemas.input_schema import InputExcelTable, NodeRead, NodePromise, ReceivedTable

SUPPORT_FILES = Path(__file__).resolve().parents[2] / "support_files" / "data"


def create_graph(flow_id: int = 1) -> FlowGraph:
    """Create a new FlowGraph for testing."""
    handler = FlowfileHandler()
    handler.register_flow(schemas.FlowSettings(flow_id=flow_id, name="test_flow", path="."))
    return handler.get_flow(flow_id)


def write_two_sheet_workbook(path: Path) -> Path:
    """A workbook whose first sheet is deliberately not the one openpyxl would pick by name."""
    workbook = Workbook()
    first = workbook.active
    first.title = "Zebra"
    first.append(["col_a"])
    first.append(["from_first_sheet"])
    second = workbook.create_sheet("Alpha")
    second.append(["col_b"])
    second.append(["from_second_sheet"])
    workbook.save(path)
    return path


def test_read_excel_table_basic():
    excel_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1"),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path(excel_table)
    except Exception as e:
        assert False, e


def test_read_excel_table_external():
    excel_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1"),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path_worker(excel_table, flow_id=-1, node_id=-1)
        flowfile_table.collect()
    except Exception as e:
        assert False, f"Failed to read excel table{e}"


def test_read_excel_table_with_type_interference():
    excel_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", type_inference=True),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path(excel_table)
        flowfile_table.collect()
    except Exception as e:
        assert False, f"Failed to read excel table{e}"


def test_read_excel_starting_from_line10():
    received_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", type_inference=True, start_row=10),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path(received_table)
        flowfile_table.collect()
    except Exception as e:
        assert False, f"Failed to read excel table{e}"


def test_read_excel_starting_from_line10_no_headers():
    received_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", type_inference=True, start_row=10, has_headers=False),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path(received_table)
        flowfile_table.count()
    except Exception as e:
        assert False, f"Failed to read excel table{e}"


def test_read_excel_starting_line_10_no_type_interference():
    received_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", start_row=10),
    )
    try:
        flowfile_table = FlowDataEngine.create_from_path(received_table)
        flowfile_table.collect()
        assert flowfile_table.columns == ['.38 Special',
                                          'Caught Up in You',
                                          'KGLK',
                                          'KGLK1446',
                                          '.38 Special_v2',
                                          'Caught Up in You by .38 Special',
                                          '0',
                                          '1402970932',
                                          'Caught Up In You']
    except:
        assert False, "Failed to read excel table"


def test_read_excel_file_second_row_date_type():
    received_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file_issue_356.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", start_row=1),
    )
    flowfile_table = FlowDataEngine.create_from_path(received_table)
    assert "2025-06-01 00:00:00" in flowfile_table.columns, f"Expected '2025-06-01 00:00:00' in columns but got {flowfile_table.columns}"


def test_read_excel_file_second_row_date_type_no_headers():
    received_table = ReceivedTable(
        path=str(SUPPORT_FILES / "excel_file_issue_356.xlsx"),
        name="excel_file",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", start_row=1, has_headers=False),
    )
    flowfile_table = FlowDataEngine.create_from_path(received_table)

    assert all(column_name.startswith("column_") for column_name in flowfile_table.columns)


def test_blank_sheet_name_normalizes_to_none():
    assert InputExcelTable(sheet_name="").sheet_name is None
    assert InputExcelTable(sheet_name="   ").sheet_name is None
    assert InputExcelTable().sheet_name is None
    assert InputExcelTable(sheet_name="Sheet1").sheet_name == "Sheet1"


def test_read_excel_without_sheet_name_reads_first_sheet(tmp_path):
    file_path = write_two_sheet_workbook(tmp_path / "two_sheets.xlsx")
    received_table = ReceivedTable(
        path=str(file_path),
        name="two_sheets.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(),
    )
    flowfile_table = FlowDataEngine.create_from_path(received_table)
    assert flowfile_table.columns == ["col_a"]
    assert flowfile_table.to_pylist() == [{"col_a": "from_first_sheet"}]


def test_read_excel_with_blank_sheet_name_reads_first_sheet(tmp_path):
    file_path = write_two_sheet_workbook(tmp_path / "two_sheets.xlsx")
    received_table = ReceivedTable(
        path=str(file_path),
        name="two_sheets.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name=""),
    )
    flowfile_table = FlowDataEngine.create_from_path(received_table)
    assert flowfile_table.columns == ["col_a"]


def test_read_excel_without_sheet_name_external(tmp_path):
    file_path = write_two_sheet_workbook(tmp_path / "two_sheets.xlsx")
    received_table = ReceivedTable(
        path=str(file_path),
        name="two_sheets.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(),
    )
    flowfile_table = FlowDataEngine.create_from_path_worker(received_table, flow_id=-1, node_id=-1)
    flowfile_table.collect()
    assert flowfile_table.columns == ["col_a"]


def test_add_read_keeps_blank_sheet_name_unresolved(tmp_path):
    """add_read must not eagerly open the workbook to pin a sheet name into the settings."""
    file_path = write_two_sheet_workbook(tmp_path / "two_sheets.xlsx")
    graph = create_graph()
    graph.add_node_promise(NodePromise(flow_id=1, node_id=1, node_type="read"))
    received_table = ReceivedTable(
        path=str(file_path),
        name="two_sheets.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name=""),
    )
    graph.add_read(NodeRead(flow_id=1, node_id=1, received_file=received_table))
    node = graph.get_node(1)
    assert node.setting_input.received_file.table_settings.sheet_name is None
    assert node.get_resulting_data().columns == ["col_a"]


def test_add_read_blank_sheet_name_does_not_open_missing_file(tmp_path):
    """A path that does not exist yet must fail at run time, not while adding the node."""
    graph = create_graph()
    graph.add_node_promise(NodePromise(flow_id=1, node_id=1, node_type="read"))
    received_table = ReceivedTable(
        path=str(tmp_path / "does_not_exist.xlsx"),
        name="does_not_exist.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name=""),
    )
    graph.add_read(NodeRead(flow_id=1, node_id=1, received_file=received_table))


def write_messy_workbook(path: Path) -> Path:
    """A sheet shaped like a real export: metadata above the table, a stray value beside it.

    The stray cell in J1 widens the worksheet's declared used range to 10 columns while the
    table itself is 8, which is what used to raise ShapeError on the calamine path.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = "Report generated by Finance"
    sheet.cell(row=1, column=10, value="stray note")
    for column, header in enumerate([f"h{i}" for i in range(8)], start=1):
        sheet.cell(row=3, column=column, value=header)
    for row in range(4, 8):
        for column in range(1, 9):
            sheet.cell(row=row, column=column, value=f"r{row}c{column}")
    workbook.save(path)
    return path


def messy_received_table(file_path: Path, **settings) -> ReceivedTable:
    return ReceivedTable(
        path=str(file_path),
        name="messy.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1", start_row=2, **settings),
    )


def test_read_messy_sheet_without_type_inference(tmp_path):
    """The reported failure: reading it raised ShapeError instead of returning the table."""
    file_path = write_messy_workbook(tmp_path / "messy.xlsx")
    flowfile_table = FlowDataEngine.create_from_path(messy_received_table(file_path))
    assert flowfile_table.columns == [f"h{i}" for i in range(8)]
    assert flowfile_table.count() == 4


def test_read_messy_sheet_with_type_inference(tmp_path):
    """The openpyxl path keeps its used-range view of the same sheet."""
    file_path = write_messy_workbook(tmp_path / "messy.xlsx")
    flowfile_table = FlowDataEngine.create_from_path(messy_received_table(file_path, type_inference=True))
    assert flowfile_table.columns[:8] == [f"h{i}" for i in range(8)]
    assert len(flowfile_table.columns) == 10


def test_read_messy_sheet_worker_matches_local(tmp_path):
    """Core and worker read the same bytes through shared.excel_reader, so they must agree."""
    file_path = write_messy_workbook(tmp_path / "messy.xlsx")
    local = FlowDataEngine.create_from_path(messy_received_table(file_path))
    remote = FlowDataEngine.create_from_path_worker(messy_received_table(file_path), flow_id=-1, node_id=-1)
    remote.collect()
    assert remote.columns == local.columns
    assert remote.to_pylist() == local.to_pylist()


def test_blank_header_cell_is_not_named_none(tmp_path):
    """A blank header cell used to become a column literally called 'None' on the calamine path."""
    file_path = tmp_path / "blank_header.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for column, header in enumerate(["a", "b", None, "d"], start=1):
        if header is not None:
            sheet.cell(row=1, column=column, value=header)
    for row in range(2, 5):
        for column in range(1, 5):
            sheet.cell(row=row, column=column, value=f"r{row}c{column}")
    workbook.save(file_path)
    received_table = ReceivedTable(
        path=str(file_path),
        name="blank_header.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name="Sheet1"),
    )
    flowfile_table = FlowDataEngine.create_from_path(received_table)
    assert flowfile_table.columns == ["a", "b", "_unnamed_column_2", "d"]


COMPLEX_FILE = SUPPORT_FILES / "complex_excel_test.xlsx"


def complex_table(sheet_name: str, **settings) -> ReceivedTable:
    """A read of one sheet of the awkward-shapes fixture (see support_files/make_complex_excel_test.py)."""
    return ReceivedTable(
        path=str(COMPLEX_FILE),
        name="complex_excel_test.xlsx",
        file_type="excel",
        table_settings=InputExcelTable(sheet_name=sheet_name, **settings),
    )


def test_complex_sales_report_ignores_the_stray_cell(tmp_path):
    """A title block above the table and a note in K1 must not widen the table to 11 columns."""
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Sales Report", start_row=4))
    assert flowfile_table.columns == ["Region", "Rep", "Q1", "Q2", "Q3", "Q4", "Total", "Margin %"]


def test_complex_sales_report_with_type_inference_keeps_cents():
    """A quarter column mixing whole and fractional amounts must widen to Float64, not fail on the float."""
    flowfile_table = FlowDataEngine.create_from_path(
        complex_table("Sales Report", start_row=4, type_inference=True)
    )
    dtypes = {name: str(dtype) for name, dtype in flowfile_table.data_frame.collect_schema().items()}
    assert dtypes["Q3"] == "Float64"
    assert 162500.25 in [row["Q3"] for row in flowfile_table.to_pylist()]


def test_complex_raw_export_names_blank_and_duplicate_headers():
    """A blank header cell and a repeated one both have to survive as usable, distinct names."""
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Raw Export"))
    assert flowfile_table.columns == [
        "order_id",
        "customer",
        "_unnamed_column_2",
        "amount",
        "amount_v2",
        "ordered_at",
        "qty",
        "active",
    ]


def test_complex_inventory_reads_despite_excel_error_values():
    """unit_cost holds floats, an int and '#N/A'; the mixed column becomes text rather than failing."""
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Inventory"))
    assert flowfile_table.columns == ["sku", "location", "on_hand", "reorder_at", "unit_cost", "status"]
    assert "#N/A" in [row["unit_cost"] for row in flowfile_table.to_pylist()]


def test_complex_timesheet_header_below_a_title_row():
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Timesheet 2026", start_row=2))
    assert flowfile_table.columns == ["employee", "mon", "tue", "wed", "thu", "fri", "notes"]


def test_complex_headers_only_sheet_yields_no_rows():
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Headers Only"))
    assert flowfile_table.columns == ["col_a", "col_b", "col_c"]
    assert flowfile_table.count() == 0


def test_complex_long_tail_reads_every_row():
    flowfile_table = FlowDataEngine.create_from_path(complex_table("Long Tail"))
    assert flowfile_table.count() == 2000


def test_complex_sales_report_matches_between_core_and_worker():
    """The fixture is the cross-service check: both sides read it through shared.excel_reader."""
    local = FlowDataEngine.create_from_path(complex_table("Sales Report", start_row=4))
    remote = FlowDataEngine.create_from_path_worker(
        complex_table("Sales Report", start_row=4), flow_id=-1, node_id=-1
    )
    remote.collect()
    assert remote.columns == local.columns
    assert remote.to_pylist() == local.to_pylist()
